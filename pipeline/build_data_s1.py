#!/usr/bin/env python3
"""Create the machine-readable Data S1 workbook for the V7 submission."""

from __future__ import annotations

import hashlib
from pathlib import Path

import pandas as pd
from openpyxl import Workbook, load_workbook
from openpyxl.cell import WriteOnlyCell
from openpyxl.styles import Font, PatternFill


V7 = Path(__file__).resolve().parents[1]
OUTPUT = V7 / "current" / "submission" / "Data_S1.xlsx"
RJ100 = V7 / "publication_figures" / "figures" / "FigS4" / "data" / "plot_data.parquet"

HEADER_FILL = PatternFill("solid", fgColor="1F4E78")
HEADER_FONT = Font(name="Arial", size=10, bold=True, color="FFFFFF")
BODY_FONT = Font(name="Arial", size=10, color="000000")


def sha256(path: Path) -> str:
    h = hashlib.sha256()
    with path.open("rb") as handle:
        for chunk in iter(lambda: handle.read(1024 * 1024), b""):
            h.update(chunk)
    return h.hexdigest()


def append_header(sheet, values) -> None:
    cells = []
    for value in values:
        cell = WriteOnlyCell(sheet, value=value)
        cell.fill = HEADER_FILL
        cell.font = HEADER_FONT
        cells.append(cell)
    sheet.append(cells)


def append_rows(sheet, rows) -> None:
    for row in rows:
        sheet.append(row)


def main() -> None:
    if not RJ100.is_file():
        raise FileNotFoundError(RJ100)
    frame = pd.read_parquet(RJ100)
    expected = [
        "panel_id", "configuration_id", "temperature_K", "Rj_nOhm", "Npw",
        "rho_turn_uOhm_cm2", "charging_time_h", "current_ramp_end_h",
        "display_end_h", "time_h", "heat_source", "temperature_stage", "heat_load_W",
    ]
    if list(frame.columns) != expected or len(frame) != 572_220:
        raise RuntimeError(f"Unexpected Rj=100 source shape/columns: {frame.shape}")

    wb = Workbook(write_only=True)

    readme = wb.create_sheet("README")
    append_header(readme, ["Field", "Value"])
    append_rows(readme, [
        ["Data supplement", "Data S1"],
        ["Title", "REBCO engineering critical-current-density values and numerical model inputs"],
        ["Purpose", "Machine-readable tables removed from the displayed Supplementary Tables, plus the numerical histories underlying the former Rj=100 nΩ charging figure."],
        ["Workbook structure", "REBCO_Jc; winding_baseline; coolant_inventory; charging_Rj100"],
        ["charging_Rj100 source", "publication_figures/figures/FigS4/data/plot_data.parquet"],
        ["charging_Rj100 source SHA-256", sha256(RJ100)],
        ["Units", "Units are encoded in column headers; each row is one observation."],
    ])
    readme.column_dimensions["A"].width = 34
    readme.column_dimensions["B"].width = 100

    jc = wb.create_sheet("REBCO_Jc")
    append_header(jc, ["B_perp_T", "J_e_4p2K_A_mm-2", "J_e_10K_A_mm-2", "J_e_20K_A_mm-2"])
    append_rows(jc, [
        [0.5, 5107, 4489, 3593], [1, 3269, 2865, 2283], [2, 2092, 1824, 1441],
        [3, 1612, 1398, 1094], [5, 1160, 995, 764], [7, 934, 792, 597],
        [10, 742, 619, 452], [15, 572, 461, 319], [20, 475, 370, 240],
        [23, 434, 330, 206],
    ])
    jc.freeze_panes = "A2"
    jc.auto_filter.ref = "A1:D11"
    for col, width in {"A": 14, "B": 20, "C": 20, "D": 20}.items():
        jc.column_dimensions[col].width = width

    baseline = wb.create_sheet("winding_baseline")
    append_header(baseline, ["parameter", "unit", "value_4p2K", "value_10K", "value_20K"])
    append_rows(baseline, [
        ["N_t", "1", 700, 850, 1150],
        ["I_p", "A", 750.0, 617.6, 456.5],
        ["L_tape", "km", 4537.6, 5510.0, 7454.7],
        ["NI_TF", "MA-turns", 8.4, 8.4, 8.4],
        ["max_eta_c", "%", 71.5, 69.8, 71.0],
    ])
    baseline.freeze_panes = "A2"
    baseline.auto_filter.ref = "A1:E6"
    for col, width in {"A": 20, "B": 14, "C": 16, "D": 16, "E": 16}.items():
        baseline.column_dimensions[col].width = width

    coolant = wb.create_sheet("coolant_inventory")
    append_header(coolant, ["coolant", "temperature_K", "pressure_bar", "density_kg_m-3", "initial_inventory_kg", "loop_volume_m3"])
    append_rows(coolant, [
        ["He", 4.2, 10, 151.150, 3507, 23.20],
        ["He", 10.0, 10, 61.273, 1422, 23.20],
        ["He", 20.0, 10, 24.244, 562, 23.20],
        ["H2", 20.0, 10, 72.405, 1680, 23.20],
    ])
    coolant.freeze_panes = "A2"
    coolant.auto_filter.ref = "A1:F5"
    for col, width in {"A": 12, "B": 18, "C": 16, "D": 20, "E": 22, "F": 18}.items():
        coolant.column_dimensions[col].width = width

    rj = wb.create_sheet("charging_Rj100")
    append_header(rj, expected)
    for row in frame.itertuples(index=False, name=None):
        rj.append(list(row))
    rj.freeze_panes = "A2"
    rj.auto_filter.ref = f"A1:M{len(frame) + 1}"
    widths = [12, 18, 15, 12, 10, 22, 18, 20, 16, 14, 18, 20, 16]
    for index, width in enumerate(widths, start=1):
        rj.column_dimensions[chr(64 + index)].width = width

    OUTPUT.parent.mkdir(parents=True, exist_ok=True)
    wb.save(OUTPUT)

    # Read-back gate: sheet names, row counts, and representative endpoint rows.
    check = load_workbook(OUTPUT, read_only=True, data_only=True)
    if check.sheetnames != ["README", "REBCO_Jc", "winding_baseline", "coolant_inventory", "charging_Rj100"]:
        raise RuntimeError(f"Unexpected sheet order: {check.sheetnames}")
    # Write-only workbooks may omit cached sheet dimensions, so read-only
    # ``max_row`` can be ``None``. Count streamed rows instead of trusting the
    # optional worksheet-dimension metadata.
    jc_rows = sum(1 for _ in check["REBCO_Jc"].iter_rows(values_only=True))
    rj_rows = sum(1 for _ in check["charging_Rj100"].iter_rows(values_only=True))
    if jc_rows != 11 or rj_rows != 572_221:
        raise RuntimeError("Data S1 row-count validation failed")
    first = tuple(check["charging_Rj100"].iter_rows(min_row=2, max_row=2, values_only=True))[0]
    last = tuple(check["charging_Rj100"].iter_rows(min_row=rj_rows, max_row=rj_rows, values_only=True))[0]
    if first[3] != 100 or last[3] != 100:
        raise RuntimeError("Data S1 Rj=100 endpoint validation failed")
    check.close()
    print(f"DATA_S1=PASS rows={len(frame)} path={OUTPUT}")


if __name__ == "__main__":
    main()
